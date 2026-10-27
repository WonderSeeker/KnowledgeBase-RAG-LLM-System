"""
知识库
"""
import os
import io
import time
import  config_data as config
import  hashlib
from langchain_chroma import Chroma
from vector_stores import CHROMA_SETTINGS
import embeddings as embeddings_module
from langchain_text_splitters import RecursiveCharacterTextSplitter
from datetime import datetime

def check_md5(md5_str:str):
    """检查传入的MD5字符串是否已经被处理过了
            return False未处理，True已处理
    """
    if not os.path.exists(config.md5_path):
        # if 进入表示文件不存在，表示没有处理过这个MD5
        open(config.md5_path,'w',encoding='utf-8').close()
        return False
    else:
        for line in open(config.md5_path,'r',encoding='utf-8').readlines():
            line=line.strip()   # 处理字符串前后的空格和回车
            if line == md5_str:
                return True     # 已处理过
        return False

def save_md5(md5_str:str):
    """将传入的md5字符串，记录到文件内保存"""
    with open(config.md5_path,'a',encoding="utf-8")as f:
        f.write(md5_str + '\n')

def get_string_md5(input_str:str ,encoding='utf-8'):
    """将出传入的字符串转换为md5字符串"""

    # 将字符串转换为bytes字节数组
    str_bytes = input_str.encode(encoding=encoding)

    # 创建md5 对象
    md5_obj =hashlib.md5()      # 得到md5对象
    md5_obj.update(str_bytes)   # 更新内容（传入即将要转换的字节数组）
    md5_hex=md5_obj.hexdigest() # 得到md5的十六进制字符串

    return md5_hex


def extract_text(file_bytes: bytes, filename: str, encoding: str = "utf-8") -> str:
    """从文件字节流中抽取纯文本

    - .pdf ：用 pypdf 逐页抽取文字（只抽文字层，扫描件抽不出来）
    - .docx：用 python-docx 抽段落 + 表格（表格按 "a | b | c" 保留，图片会丢）
    - 其它 ：按文本解码，先试 utf-8，失败再试 gbk（兼容中文 Windows 的 txt）

    :param file_bytes: 文件原始字节
    :param filename: 文件名（用于判断类型）
    :param encoding: 文本文件的首选编码
    """
    lower = filename.lower()

    if lower.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)

    # 改造：补 DOCX 支持（Word 2007+）。老的 .doc 二进制格式 python-docx 读不了，
    # 需要在 Word 里另存为 .docx 或 .txt。
    if lower.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:                      # 表格按行拼成 "a | b | c"
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    if lower.endswith(".doc"):
        raise ValueError("不支持旧版 .doc，请先另存为 .docx 或 .txt")

    try:
        return file_bytes.decode(encoding)
    except UnicodeDecodeError:
        return file_bytes.decode("gbk", errors="ignore")


class KnowledgeBaseService(object):
    def __init__(self, chunk_size=None, chunk_overlap=None, collection_name=None,
                 embedding_backend=None):
        """可以覆盖分块参数、集合名与嵌入后端

        :param chunk_size: 覆盖 config.chunk_size
        :param chunk_overlap: 覆盖 config.chunk_overlap
        :param collection_name: 覆盖默认集合名（默认按嵌入后端区分：rag / rag_local）
        :param embedding_backend: dashscope | local，覆盖 config.embedding_backend
        """
        self.chunk_size = config.chunk_size if chunk_size is None else chunk_size
        self.chunk_overlap = config.chunk_overlap if chunk_overlap is None else chunk_overlap

        # 嵌入后端：解析 profile → local 探活 → 实测真实维度（唯一权威来源）
        self.profile, self.embeddings, self.embedding_dim = embeddings_module.prepare(embedding_backend)
        self.collection_name = collection_name or embeddings_module.default_collection_name(embedding_backend)

        # 如果文件夹不存在则创建，如果存在则跳过
        os.makedirs(config.persist_directory,exist_ok=True)

        self.chroma=Chroma(          # 向量存储的示例 Chroma向量库对象
            collection_name=self.collection_name,        #数据库表名
            # 改造：原来这里硬编码 "text-embedding-v4"，与 rag.py 走 config 的写法不一致，
            # 一旦改配置就会出现「入库用 A 模型、检索用 B 模型」导致检索失效
            embedding_function=self.embeddings,
            persist_directory=config.persist_directory,   #数据库本地存储文件夹
            client_settings=CHROMA_SETTINGS,              # 关闭匿名遥测，避免脏日志
        )      # 向量存储的实例，Chroma向量库对象
        # 指纹校验：集合里的向量必须与当前嵌入后端/维度一致，否则直接报错退出
        embeddings_module.verify_and_stamp(
            self.chroma._collection, self.profile, self.embedding_dim, self.collection_name
        )
        self.spliter=RecursiveCharacterTextSplitter(  # 文本分割器的对象
            chunk_size=self.chunk_size,               # 分割后的文本段最大长度
            chunk_overlap=self.chunk_overlap,         # 连续文本段之间的字符重叠数量
            separators=config.separators,             # 自然段落划分的符号
            length_function=len,                      # 使用python自带的len函数做长度统计的依赖
        )      # 文本分割器的对象

    def upload_by_str(self,data:str,filename):
        """将传入的字符串，进行向量化，存入向量数据库中"""
        # 先得到出传入的字符串的md5值
        md5_hex=get_string_md5(data)
        if check_md5(md5_hex):
            return "[Repeat] 内容已存在知识库"
        # 缺陷修复（阶段 3 补充）：原来用 max_spliter_char_number(=1000) 做硬门槛，
        #   if len(data) > 1000: 切分 else: [data]
        # 导致 ≤1000 字的文档「永远不切分」，《尺码推荐.txt》(270 字) 无论 chunk_size
        # 设 300 还是 1500 都只有 1 个片段 —— 分块参数对短文档完全失效。
        # 现在一律交给 splitter：文本短于 chunk_size 时它自然返回 1 个片段，
        # 文本更长时按 chunk_size/chunk_overlap 切分。判定语义单一，chunk_size 真实生效。
        knowledge_chunks:list[str]=self.spliter.split_text(data)

        # 基线观测：记录下来，方便核对分块参数的实际效果
        print(f"[入库] 文件={filename} 原文长度={len(data)} 切分片段数={len(knowledge_chunks)} "
              f"(chunk_size={self.chunk_size} chunk_overlap={self.chunk_overlap} "
              f"集合={self.collection_name} 嵌入={self.profile.backend}/{self.embedding_dim}维)")

        base_metadata={
            "source":filename,
            #2026-3-8 15:43:30
            "create_time":datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "operator":"客户",
        }

        start_time=time.perf_counter()   # 基线观测：记录向量化 + 写库耗时
        # 缺陷修复（上游遗留，静默型）：LangChain 的参数名是 metadatas（复数），
        # 原项目写成 metadata（单数）→ 落进 **kwargs 被静默忽略 → 所有片段元数据全丢
        # （实测 6/6 片段的 metadata 都是 None）。
        # 后果：不知道片段来自哪个文档、format_document 拼给模型的是空元数据、
        #       文档级归因（IR 指标 / 分块实验的按文档统计）完全做不了。
        self.chroma.add_texts(        # 内容加载到向量库中
            # iterable-> list \tuple
            knowledge_chunks,
            metadatas=[{**base_metadata, "chunk_index": index}
                       for index, _ in enumerate(knowledge_chunks)],
        )
        cost=time.perf_counter()-start_time
        save_md5(md5_hex)
        print(f"[入库] 向量化+写库耗时 {cost:.2f}s")
        return f"[Success]内容已经成功载入向量库（{len(knowledge_chunks)} 个片段，耗时 {cost:.2f}s）"

if __name__ =='__main__':
    # r1 = get_string_md5("周杰伦")
    # r2 = get_string_md5("周杰伦")
    # r3 = get_string_md5("周杰伦4")
    # # 为什么用md5  哈希算法（散列函数）能把任意长度数据，算出一串32位16进制的字符串
    # #优点： 快 简单 能做简单文件校验   ，缺点：  不能做安全加密   易于被破解，过时
    # #SHA-256  SHA-512
    #
    # print(r1)
    # print(r2)
    # print(r3)

    #save_md5("7a8941058aaf4df5147042ce104568da")
    #print(check_md5("7a8941058aaf4df5147042ce104568da"))
    service= KnowledgeBaseService()
    r=service.upload_by_str("流星","testfile")
    print(r)
